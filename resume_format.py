from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json
from datetime import datetime

def create_resume(data):
    # Create a new Document
    doc = Document()
    
    # Set margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
    
    # Header with name and contact info
    header = doc.add_paragraph()
    name_run = header.add_run(data["name"])
    name_run.bold = True
    name_run.font.size = Pt(18)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Contact info line
    contact_info = []
    if data.get('email'):
        contact_info.append(data['email'])
    if data.get('phone'):
        contact_info.append(data['phone'])
    if data.get('location'):
        contact_info.append(data['location'])
    
    # Add social links
    social_links = []
    for profile in data.get('social_profiles', []):
        social_links.append(f"{profile['platform']}: {profile['url']}")
    
    contact_line = doc.add_paragraph()
    contact_line.add_run(" | ".join(contact_info + social_links))
    contact_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add a professional summary
    if data.get('description'):
        doc.add_heading('Professional Summary', level=1)
        doc.add_paragraph(data['description'])
    
    # Experience Section
    if data.get('experiences') and len(data['experiences']) > 0:
        doc.add_heading('Professional Experience', level=1)
        for exp in data['experiences']:
            p = doc.add_paragraph()
            p.add_run(f"{exp['title']} - {exp['company']}").bold = True
            p.add_run(f"\n{exp['duration']}")
            doc.add_paragraph(exp['description']).style = 'List Bullet'
    
    # Education Section
    if data.get('education') and len(data['education']) > 0:
        doc.add_heading('Education', level=1)
        for edu in data['education']:
            p = doc.add_paragraph()
            degree_line = f"{edu['degree']}, {edu['institution']} ({edu['year']})"
            if edu.get('cgpa'):
                degree_line += f" - CGPA: {edu['cgpa']}"
            p.add_run(degree_line).bold = True
    
    # Skills Section
    if data.get('skills') and len(data['skills']) > 0:
        doc.add_heading('Skills', level=1)
        
        # Group skills by level
        skill_levels = {"Advanced": [], "Intermediate": [], "Beginner": []}
        for skill in data['skills']:
            skill_levels[skill['level']].append(skill['name'])
        
        for level, skills in skill_levels.items():
            if skills:
                p = doc.add_paragraph()
                p.add_run(f"{level}: ").bold = True
                p.add_run(", ".join(skills))
    
    # Projects Section
    if data.get('projects') and len(data['projects']) > 0:
        doc.add_heading('Projects', level=1)
        for project in data['projects']:
            p = doc.add_paragraph()
            p.add_run(f"{project['title']}").bold = True
            p.add_run(f" ({', '.join(project['tools'])})")
            doc.add_paragraph(project['description']).style = 'List Bullet'
    
    # Certifications Section
    if data.get('certifications') and len(data['certifications']) > 0:
        doc.add_heading('Certifications', level=1)
        for cert in data['certifications']:
            p = doc.add_paragraph()
            p.add_run(f"{cert['name']} - {cert['issuer']} ({cert['date']})").italic = True
    
    # Additional Sections
    if data.get('achievements') and len(data['achievements']) > 0:
        doc.add_heading('Achievements', level=1)
        for achievement in data['achievements']:
            doc.add_paragraph(achievement, style='List Bullet')
    
    if data.get('languages') and len(data['languages']) > 0:
        doc.add_heading('Languages', level=1)
        doc.add_paragraph(", ".join(data['languages']))
    
    if data.get('interests') and len(data['interests']) > 0:
        doc.add_heading('Interests', level=1)
        doc.add_paragraph(", ".join(data['interests']))
    
    # Format headings for ATS compatibility
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith('Heading'):
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(14)
    
    # Add metadata for ATS parsing
    core_properties = doc.core_properties
    core_properties.title = f"Resume - {data['name']}"
    core_properties.subject = f"Resume | {data.get('specialization', '')}"
    core_properties.keywords = ", ".join([skill['name'] for skill in data.get('skills', [])])
    core_properties.language = "en-US"
    core_properties.created = datetime.now()
    
    return doc

def load_json_data(json_path):
    with open(json_path, 'r') as file:
        return json.load(file)



if __name__ == "__main__":
    file = load_json_data("AI_Enhanced_CV.json")
    resume = create_resume(file)
    resume.save("./output/trail1.docx")